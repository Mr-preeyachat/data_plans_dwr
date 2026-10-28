import os
import sys
import json
import time
import datetime
import threading
import urllib.request
import csv
import io
import webbrowser
from http.server import HTTPServer, SimpleHTTPRequestHandler
from urllib.parse import parse_qs, urlparse

# Force UTF-8 stdout/stderr on Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

PORT = 8000
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS_DIR = os.path.join(BASE_DIR, "assets")
DATA_JSON_PATH = os.path.join(ASSETS_DIR, "data.json")

# Google Sheets Configuration
GOOGLE_SHEET_ID = "18bZNi3IlQNiUHtVhUvsZV65zA2EpqJRyl1ARgDRKsN0"
GOOGLE_SHEET_NAME = "DWR_007"
# NOTE: ใช้ export?format=xlsx แทน gviz CSV เพราะ gviz จะตัดค่าข้อความทิ้ง (กลายเป็นว่าง)
# เมื่อคอลัมน์มีข้อมูลปนกัน (เช่น คอลัมน์ Q: 2571 ปนกับ "2567 (งบกลาง 1)")
GOOGLE_SHEET_XLSX_URL = f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}/export?format=xlsx"
START_ROW = 6  # ข้อมูลเริ่มที่แถวที่ 6
LOCAL_EXCEL_NAMES = ["DWR007.xlsx", "แบบฟอร์ม_DWR007_260820.xlsm"]
SOURCE_GOOGLE = "Google Sheets (Live Sync)"
SOURCE_EXCEL = "Excel สำรอง (ฉุกเฉิน)"
ENABLE_GOOGLE_SHEETS = True

# Column Map: Excel Column (1-indexed) -> Field Key
COL_MAP = {
    4: "D",   # หน่วยงาน
    5: "E",   # ชื่อโครงการ
    6: "F",   # หมู่บ้าน
    7: "G",   # หมู่ที่
    8: "H",   # ตำบล
    9: "I",   # อำเภอ
    10: "J",  # จังหวัด
    11: "K",  # ภาค
    12: "L",  # ลุ่มน้ำหลัก
    13: "M",  # ลุ่มน้ำสาขา
    14: "lat",# ละติจูด
    15: "lon",# ลองจิจูด
    16: "P",  # งบประมาณ
    17: "Q",  # ประเภทงบประมาณ
    24: "X",  # สนับสนุน
    29: "AC", # ประเภทแผนงาน
    30: "AD", # กิจกรรม-ลักษณะงาน
    31: "AE", # ประเภทโครงการ
    33: "AG", # ปริมาณน้ำเพิ่มขึ้น
    34: "AH", # ครัวเรือนได้รับประโยชน์
    35: "AI", # พื้นที่การเกษตรได้รับประโยชน์
    36: "AJ", # ปริมาณน้ำที่กระจายได้
    37: "AK"  # ระยะทาง
}

NUMERIC_FIELDS = {"P", "AG", "AH", "AI", "AJ", "AK", "lat", "lon"}

cached_data = []
last_mtime = 0
last_filename = ""
data_source = ""
last_fetch_time = 0
CACHE_TTL_SECONDS = 30  # Re-fetch Google Sheets at most every 30 seconds

def clean_num(val):
    if val is None or val == "":
        return 0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0

def to_text(val):
    """แปลงค่าเซลล์เป็นข้อความ (2571.0 -> '2571', ข้อความคงเดิม)"""
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val).replace("\n", " ").strip()
    while "  " in s:
        s = s.replace("  ", " ")
    return s

def row_to_item(row):
    item = {}
    for col_1idx, key in COL_MAP.items():
        val_idx = col_1idx - 1
        val = row[val_idx] if val_idx < len(row) else None
        item[key] = clean_num(val) if key in NUMERIC_FIELDS else to_text(val)
    return item

def parse_workbook(source):
    """อ่านไฟล์ xlsx (path หรือ BytesIO) -> list ของโครงการ"""
    import openpyxl
    wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    rows_data = []
    try:
        sheet_name = GOOGLE_SHEET_NAME if GOOGLE_SHEET_NAME in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if idx < START_ROW or not row or len(row) < 5:
                continue
            if not to_text(row[4]):  # ไม่มีชื่อโครงการ -> ข้าม
                continue
            rows_data.append(row_to_item(row))
    finally:
        wb.close()
    return rows_data

def fetch_google_sheet():
    print(f"[INFO] Fetching live data from Google Sheets ({GOOGLE_SHEET_NAME}) ...")
    start_time = time.time()
    try:
        req = urllib.request.Request(GOOGLE_SHEET_XLSX_URL, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=60) as res:
            raw_bytes = res.read()
        if not raw_bytes.startswith(b"PK"):
            raise ValueError("ไม่ได้รับไฟล์ xlsx (ตรวจสอบว่าแชร์ชีตเป็น 'ทุกคนที่มีลิงก์ดูได้')")
        rows_data = parse_workbook(io.BytesIO(raw_bytes))
        print(f"[INFO] Successfully parsed {len(rows_data)} rows from Google Sheets in {time.time()-start_time:.2f}s")
        return rows_data
    except Exception as e:
        print(f"[WARNING] Failed to fetch Google Sheet: {e}")
        return None

def find_excel_file():
    """หาไฟล์ Excel สำรอง: ข้างๆ server.py ก่อน แล้วค่อยดูในโฟลเดอร์ assets/"""
    for folder in (BASE_DIR, ASSETS_DIR):
        if not os.path.isdir(folder):
            continue
        for filename in LOCAL_EXCEL_NAMES:
            filepath = os.path.join(folder, filename)
            if os.path.exists(filepath):
                return filepath
    for folder in (BASE_DIR, ASSETS_DIR):
        if not os.path.isdir(folder):
            continue
        for filename in sorted(os.listdir(folder)):
            if filename.lower().endswith((".xlsx", ".xlsm")) and not filename.startswith("~$"):
                return os.path.join(folder, filename)
    return None

def parse_excel():
    global cached_data, last_mtime, last_filename, data_source
    excel_path = find_excel_file()
    if not excel_path:
        print("[WARNING] ไม่พบไฟล์ Excel สำรอง (วาง DWR007.xlsx ไว้ข้าง server.py)")
        return []

    mtime = os.path.getmtime(excel_path)
    filename = os.path.basename(excel_path)
    print(f"[INFO] Reading local Excel file: {filename} ...")
    start_time = time.time()
    try:
        rows_data = parse_workbook(excel_path)
    except Exception as e:
        print(f"[ERROR] Failed reading Excel (ต้องติดตั้ง openpyxl: pip install openpyxl): {e}")
        return []

    print(f"[INFO] Parsed {len(rows_data)} rows from Excel in {time.time()-start_time:.2f}s")
    last_mtime = mtime
    last_filename = filename
    data_source = SOURCE_EXCEL
    save_data_json(rows_data, filename, data_source, mtime)
    return rows_data

def save_data_json(rows_data, source_name, source_type, mtime=None):
    try:
        os.makedirs(ASSETS_DIR, exist_ok=True)
        mtime_str = datetime.datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M:%S") if mtime else datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        payload = {
            "status": "success",
            "source": source_type,
            "file": source_name,
            "updated_at": mtime_str,
            "total": len(rows_data),
            "data": rows_data
        }
        with open(DATA_JSON_PATH, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"[WARNING] Could not save data.json: {e}")

def get_data(force_refresh=False):
    global cached_data, last_mtime, last_filename, data_source, last_fetch_time
    now = time.time()

    # Cache TTL check
    if not force_refresh and cached_data and (now - last_fetch_time < CACHE_TTL_SECONDS):
        return cached_data

    # Attempt 1: Fetch Google Sheets
    if ENABLE_GOOGLE_SHEETS:
        gs_data = fetch_google_sheet()
        if gs_data is not None and len(gs_data) > 0:
            cached_data = gs_data
            data_source = SOURCE_GOOGLE
            last_filename = f"Google Sheet ({GOOGLE_SHEET_NAME})"
            last_mtime = now
            last_fetch_time = now
            save_data_json(gs_data, last_filename, data_source)
            return cached_data

    # Attempt 2: Fallback to local Excel file
    excel_data = parse_excel()
    if excel_data:
        cached_data = excel_data
        last_fetch_time = now
        return cached_data

    return cached_data

class DashboardRequestHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def end_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/data":
            query = parse_qs(parsed.query)
            force = "force" in query or "refresh" in query
            data = get_data(force_refresh=force)
            mtime_str = datetime.datetime.fromtimestamp(last_mtime).strftime("%Y-%m-%d %H:%M:%S") if last_mtime else ""
            payload = {
                "status": "success",
                "source": data_source,
                "file": last_filename,
                "updated_at": mtime_str,
                "total": len(data),
                "data": data
            }
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        if parsed.path == "/api/status":
            mtime_str = datetime.datetime.fromtimestamp(last_mtime).strftime("%Y-%m-%d %H:%M:%S") if last_mtime else ""
            payload = {
                "status": "ok",
                "source": data_source,
                "file": last_filename,
                "last_modified": mtime_str,
                "total": len(cached_data)
            }
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        super().do_GET()

def open_browser():
    time.sleep(1.2)
    print(f"[INFO] Opening browser at http://localhost:{PORT}")
    webbrowser.open(f"http://localhost:{PORT}")

def main():
    if "--test" in sys.argv:
        print("[TEST] Running Data parser test...")
        data = get_data(force_refresh=True)
        print(f"[TEST] Source: {data_source}")
        print(f"[TEST] Total rows parsed: {len(data)}")
        if data:
            print("[TEST] Row 1 Sample:", json.dumps(data[0], ensure_ascii=False, indent=2))
        return

    print("=" * 60)
    print(" DWR007 Project Dashboard Server")
    print(" Google Sheets Sync & Local Excel Fallback")
    print("=" * 60)
    
    # Pre-parse data on startup
    get_data(force_refresh=True)

    server = HTTPServer(("0.0.0.0", PORT), DashboardRequestHandler)
    print(f"[SUCCESS] Server running at http://localhost:{PORT}")
    print("[INFO] Press Ctrl+C to stop server")

    # Auto open browser in thread
    threading.Thread(target=open_browser, daemon=True).start()

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n[INFO] Server stopped.")

if __name__ == "__main__":
    main()
